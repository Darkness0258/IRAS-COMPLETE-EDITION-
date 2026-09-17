from __future__ import annotations

from pathlib import Path
import csv
import json
import zipfile
from typing import Any


class ArtifactEngine:
    def __init__(self, root: str | Path):
        self.root = Path(root)
        self.root.mkdir(parents=True, exist_ok=True)

    def _path(self, name: str, suffix: str) -> Path:
        safe = "".join(c for c in name if c.isalnum() or c in "-_. ").strip().replace(" ", "_") or "artifact"
        if not safe.lower().endswith(suffix): safe += suffix
        path = (self.root / safe).resolve()
        if self.root.resolve() not in path.parents:
            raise ValueError("Artifact path escaped root.")
        return path

    def document(self, name: str, paragraphs: list[str]) -> str:
        from docx import Document
        path = self._path(name, ".docx")
        doc = Document()
        for p in paragraphs: doc.add_paragraph(str(p))
        doc.save(path)
        return str(path)

    def pdf(self, name: str, lines: list[str]) -> str:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        path = self._path(name, ".pdf")
        c = canvas.Canvas(str(path), pagesize=A4); width,height=A4; y=height-50
        for raw in lines:
            for line in str(raw).splitlines() or [""]:
                if y < 50: c.showPage(); y=height-50
                c.drawString(50,y,line[:110]); y-=16
        c.save(); return str(path)

    def spreadsheet(self, name: str, rows: list[list[Any]]) -> str:
        from openpyxl import Workbook
        path = self._path(name, ".xlsx"); wb=Workbook(); ws=wb.active
        for row in rows: ws.append(list(row))
        wb.save(path); return str(path)

    def presentation(self, name: str, slides: list[dict[str, Any]]) -> str:
        from pptx import Presentation
        path = self._path(name, ".pptx"); prs=Presentation()
        for item in slides:
            slide=prs.slides.add_slide(prs.slide_layouts[1]); slide.shapes.title.text=str(item.get("title") or "")
            slide.placeholders[1].text=str(item.get("body") or "")
        prs.save(path); return str(path)

    def bundle(self, name: str, files: list[str | Path]) -> str:
        path=self._path(name,".zip")
        with zipfile.ZipFile(path,"w",zipfile.ZIP_DEFLATED) as z:
            for file in files:
                p=Path(file)
                if p.is_file(): z.write(p, arcname=p.name)
        return str(path)
